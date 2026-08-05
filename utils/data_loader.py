# -*- coding: utf-8 -*-
"""
数据驱动加载器（data_loader）
===================================================================
职责：
1. load_yaml(path)        —— 加载 YAML 用例数据
2. load_excel(path, sheet)—— 加载 Excel 用例数据（openpyxl），首行为表头
3. replace_variables()    —— 递归替换用例中的 ${var} 变量（支持 dict/list/str）
4. build_param_ids()      —— 生成 pytest.mark.parametrize 友好的用例 ID

设计要点：
- 数据结构统一为「list[dict]」，每条用例一个字典，可直接喂给 parametrize
- 变量替换支持 ${var} 与 ${env.VAR}，变量来源优先级：
  显式传入 variables > 环境变量 os.environ
- 支持 ${var:default} 默认值语法
"""
import os
import re
from pathlib import Path
from typing import Any, List, Dict

import yaml

from utils.logger import get_logger

log = get_logger("data")

# 匹配 ${var} / ${var:default} / ${env.VAR}
_VAR_PATTERN = re.compile(r"\$\{([a-zA-Z_][\w\.]*)(?::([^}]*))?\}")


def load_yaml(path) -> List[Dict[str, Any]]:
    """
    加载 YAML 用例文件，返回用例列表

    :param path: YAML 文件路径
    :return: list[dict]，每条用例一个字典
    """
    path = Path(path)
    if not path.exists():
        log.error("YAML 用例文件不存在: {}", path)
        raise FileNotFoundError(path)

    with open(path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)

    # 统一成 list[dict]：单个 dict 也包装成列表
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        raise ValueError(f"YAML 用例文件 [{path}] 根节点应为 list 或 dict")

    log.debug("加载 YAML 用例: {} | 共 {} 条", path.name, len(data))
    return data


def load_excel(path, sheet: str = None) -> List[Dict[str, Any]]:
    """
    加载 Excel 用例文件，首行作为表头，每行转为一条用例字典

    :param path:  Excel 文件路径
    :param sheet: 工作表名，默认第一个表
    :return: list[dict]
    """
    import openpyxl  # 延迟导入，非 YAML 用例场景无需强依赖

    path = Path(path)
    if not path.exists():
        log.error("Excel 用例文件不存在: {}", path)
        raise FileNotFoundError(path)

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration:
        log.error("Excel 文件 [{}}] 无数据行", path.name)
        return []

    cases: List[Dict[str, Any]] = []
    for idx, row in enumerate(rows, start=2):  # 数据从第 2 行开始
        if all(v is None for v in row):  # 跳过空行
            continue
        case = {}
        for key, val in zip(headers, row):
            if not key:
                continue
            # 尝试把纯数字字符串还原为 int/float，便于断言
            case[key] = _auto_cast(val)
        # 补充行号，便于失败定位
        case["_row"] = idx
        cases.append(case)

    wb.close()
    log.debug("加载 Excel 用例: {}[{}] | 共 {} 条", path.name, ws.title, len(cases))
    return cases


def _auto_cast(val):
    """把 Excel 单元格值做轻量类型推断"""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    val = str(val).strip()
    if val == "":
        return ""
    # 尝试 int / float
    try:
        if "." not in val:
            return int(val)
        return float(val)
    except ValueError:
        return val


def replace_variables(data: Any, variables: Dict[str, Any] = None) -> Any:
    """
    递归替换数据结构中的 ${var} 变量

    :param data:     dict / list / str，待替换的数据
    :param variables: 变量字典，优先级高于环境变量
    :return: 替换后的数据（同结构）

    示例：
        >>> replace_variables({"url": "${base_url}/login"}, {"base_url": "http://a.com"})
        {'url': 'http://a.com/login'}
    """
    variables = variables or {}

    if isinstance(data, dict):
        return {k: replace_variables(v, variables) for k, v in data.items()}
    if isinstance(data, list):
        return [replace_variables(v, variables) for v in data]
    if isinstance(data, str):
        return _replace_str(data, variables)
    return data


def _replace_str(text: str, variables: Dict[str, Any]) -> Any:
    """替换字符串中的 ${var} / ${var:default} / ${env.VAR}"""

    def _sub(match):
        var_name = match.group(1)
        default = match.group(2)  # 可能为 None

        # 1) 显式 variables 字典
        if var_name in variables:
            return str(variables[var_name])

        # 2) env.VAR 形式 → 取环境变量
        if var_name.startswith("env."):
            env_key = var_name[4:]
            return os.getenv(env_key, default if default is not None else match.group(0))

        # 3) 普通环境变量
        if var_name in os.environ:
            return os.environ[var_name]

        # 4) 默认值
        if default is not None:
            return default

        # 5) 找不到变量，保留原样并告警
        log.warning("未找到变量 ${{{}}}, 已保留原样", var_name)
        return match.group(0)

    replaced = _VAR_PATTERN.sub(_sub, text)

    # 若整串就是一个变量引用，尝试还原为原始类型（int/bool/None）
    full = _VAR_PATTERN.fullmatch(text)
    if full:
        return _restore_type(replaced)
    return replaced


def _restore_type(val: str):
    """把 'true'/'false'/'null'/'123' 还原为 Python 类型"""
    if val == "true":
        return True
    if val == "false":
        return False
    if val in ("null", "None"):
        return None
    try:
        return int(val)
    except ValueError:
        pass
    try:
        return float(val)
    except ValueError:
        return val


def build_param_ids(cases: List[Dict[str, Any]], key: str = "case_name") -> List[str]:
    """
    生成 pytest.mark.parametrize 的 ids 列表

    :param cases: 用例列表
    :param key:   用例名字段，默认 case_name
    :return: ids 列表
    """
    ids = []
    for idx, case in enumerate(cases, start=1):
        name = case.get(key) or case.get("title") or f"case_{idx}"
        # 清理非法字符，避免 pytest ids 报错
        safe = re.sub(r"[\\/\s\[\]:]", "_", str(name))
        ids.append(f"{idx}.{safe}")
    return ids
